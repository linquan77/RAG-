from langchain_community.document_loaders import PyMuPDFLoader, Docx2txtLoader, WebBaseLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import Chroma
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_core.documents import Document
import openpyxl
import os
from config import *


def get_embeddings():
    return HuggingFaceEmbeddings(
        model_name=EMBEDDING_MODEL,
        model_kwargs={"device": "cpu"},
        encode_kwargs={"normalize_embeddings": True}
    )


def load_excel(file_path: str):
    """
    解析 Excel 价格表：
    - 第一行为表头（序号、商品名称、单价等）
    - 单价为空表示暂无价格
    - 每行转成一段自然语言描述存入向量库
    """
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    # 第一行是表头，读取列名
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    # 找到商品名称列和单价列的索引
    name_col = None
    price_col = None
    for i, h in enumerate(headers):
        if "名称" in h or "商品" in h or "品名" in h:
            name_col = i
        if "单价" in h or "价格" in h or "价" in h:
            price_col = i

    # 找不到列名则默认第二列是名称，第三列是单价
    if name_col is None:
        name_col = 1
    if price_col is None:
        price_col = 2

    docs = []
    file_name = os.path.basename(file_path)
    all_items = []

    for row_idx, row in enumerate(rows[1:], start=2):  # 跳过表头
        if not any(row):  # 跳过空行
            continue

        name = row[name_col] if name_col < len(row) else None
        price = row[price_col] if price_col < len(row) else None

        if name is None or str(name).strip() == "":
            continue

        name = str(name).strip()

        # 生成自然语言描述
        if price is None or str(price).strip() == "":
            text = f"{name}：暂无价格信息。"
        else:
            try:
                price_float = float(price)
                price_str = str(int(price_float)) if price_float == int(price_float) else str(round(price_float, 2))
            except (ValueError, TypeError):
                price_str = str(price).strip()
            text = f"{name}的单价是{price_str}元。"

        all_items.append(name)
        docs.append(Document(
            page_content=text,
            metadata={"source": file_name, "row": row_idx, "type": "price_table"}
        ))

    # 加一个汇总文档，方便回答"有哪些商品"之类的问题
    if all_items:
        summary = f"价格表【{file_name}】共包含以下商品：\n" + "、".join(all_items[:100])
        if len(all_items) > 100:
            summary += f"……等共 {len(all_items)} 种商品。"
        docs.insert(0, Document(
            page_content=summary,
            metadata={"source": file_name, "type": "price_summary"}
        ))

    return docs


def load_file(file_path: str):
    ext = file_path.split(".")[-1].lower()
    if ext == "pdf":
        return PyMuPDFLoader(file_path).load()
    elif ext == "docx":
        return Docx2txtLoader(file_path).load()
    elif ext in ["xlsx", "xls"]:
        return load_excel(file_path)
    elif file_path.startswith("http"):
        return WebBaseLoader(file_path).load()
    else:
        raise ValueError(f"不支持的格式: {ext}")


def ingest(file_path: str, original_name: str = None, progress_callback=None):
    # 1. 加载文档
    if progress_callback:
        progress_callback(0.1, "加载文件...")
    docs = load_file(file_path)

    # 2. 设置 source 元数据为原始文件名
    display_name = original_name or file_path
    for doc in docs:
        doc.metadata["source"] = display_name

    # 3. 切块（Excel 每行已是独立语义单元，不再切块）
    if progress_callback:
        progress_callback(0.3, "分割内容...")
    ext = file_path.split(".")[-1].lower()
    if ext in ["xlsx", "xls"]:
        chunks = docs
    else:
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=CHUNK_SIZE,
            chunk_overlap=CHUNK_OVERLAP
        )
        chunks = splitter.split_documents(docs)

    # 4. 去重检查
    if progress_callback:
        progress_callback(0.5, "检查重复...")
    vectorstore = Chroma(
        persist_directory=CHROMA_DB_PATH,
        embedding_function=get_embeddings()
    )

    existing = vectorstore.get(where={"source": display_name})
    if existing and len(existing["ids"]) > 0:
        if progress_callback:
            progress_callback(1.0, "已跳过")
        return 0

    # 5. 写入向量库
    if progress_callback:
        progress_callback(0.7, "写入数据库...")
    vectorstore.add_documents(chunks)

    if progress_callback:
        progress_callback(1.0, "完成！")
    return len(chunks)